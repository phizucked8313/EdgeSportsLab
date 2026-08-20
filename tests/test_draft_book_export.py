import pandas as pd
from openpyxl import load_workbook

from fantasy_draft_model.exports.csv_export import write_csv_tables
from fantasy_draft_model.exports.draft_book_export import build_export_tables
from fantasy_draft_model.exports.draft_book_export import write_pdf_draft_book
from fantasy_draft_model.exports.excel_export import write_excel_workbook
from fantasy_draft_model.build_draft_book import build_draft_package
from fantasy_draft_model.models.league_profile import get_league


def _rankings():
    positions = ["QB", "RB", "WR", "TE", "K", "DEF"]
    rows = []
    for rank in range(1, 306):
        position = positions[(rank - 1) % len(positions)]
        rows.append(
            {
                "draft_rank": rank,
                "player_name_clean": f"Player {rank}",
                "position": position,
                "position_rank_label": f"{position}{(rank - 1) // 6 + 1}",
                "team": "BUF",
                "bye_week": 7,
                "tier": (rank - 1) // 24 + 1,
                "projected_points": 350.0 - rank,
                "vorp": 100.0 - rank,
                "edgescore": 90.0 - rank / 10,
                "draft_score": 95.0 - rank / 10,
                "injury_risk_score": 10.0,
                "injury_risk_label": "LOW",
                "current_injury_status": "Questionable" if rank == 3 else "",
                "current_injury_body_part": "Knee" if rank == 3 else "",
                "current_injury_data_quality": "A" if rank == 3 else "",
                "current_injury_source": "fixture" if rank == 3 else "",
            }
        )
    return pd.DataFrame(rows)


def test_export_tables_preserve_rankings_and_annotate_keepers_and_user_picks():
    rankings = _rankings()
    original = rankings.copy(deep=True)
    keepers = pd.DataFrame(
        [
            {
                "owner_team": "BLKWDW'S",
                "player_name": "Player 2",
                "keeper_type": "standard",
                "keeper_round": 15,
            }
        ]
    )

    tables = build_export_tables(rankings, keepers, get_league("Drunk Sundays"))

    pd.testing.assert_frame_equal(rankings, original)
    assert list(tables) == ["Top 300", "QB", "RB", "WR", "TE", "K", "DEF", "BLKWDW Picks"]
    assert tables["Top 300"]["Overall Rank"].tolist() == list(range(1, 301))
    assert tables["Top 300"]["Draft Score"].tolist() == original.head(300)["draft_score"].tolist()
    keeper = tables["Top 300"].loc[tables["Top 300"]["Player"] == "Player 2"].iloc[0]
    assert keeper["Availability"] == "UNAVAILABLE - KEEPER"
    assert keeper["Keeper"] == "BLKWDW'S (R15 standard)"
    assert tables["Top 300"].loc[8, "BLKWDW Pick"] == "PICK 9 (R1)"
    assert tables["Top 300"].loc[15, "BLKWDW Pick"] == "PICK 16 (R2)"
    assert tables["BLKWDW Picks"]["Overall Pick"].tolist()[:3] == [9, 16, 33]
    assert len(tables["QB"]) == 51
    assert tables["QB"]["Position Tier"].map(lambda value: isinstance(value, int)).all()
    assert tables["Top 300"].loc[2, "Current Injury"] == "Questionable - Knee"


def test_user_keeper_rounds_replace_live_pick_markers():
    keepers = pd.DataFrame(
        [
            {
                "owner_team": "BLKWDW'S",
                "player_name": "Ashton Jeanty",
                "keeper_type": "rookie",
                "keeper_round": 3,
            },
            {
                "owner_team": "BLKWDW'S",
                "player_name": "TreVeyon Henderson",
                "keeper_type": "standard",
                "keeper_round": 15,
            },
            {
                "owner_team": "Other Team",
                "player_name": "Someone Else",
                "keeper_type": "standard",
                "keeper_round": 1,
            },
        ]
    )

    tables = build_export_tables(_rankings(), keepers, get_league("Drunk Sundays"))

    schedule = tables["BLKWDW Picks"].set_index("Round")
    assert schedule.loc[1, "Marker"] == "PICK 9 (R1)"
    assert schedule.loc[3, "Marker"] == "KEEPER: Ashton Jeanty"
    assert schedule.loc[15, "Marker"] == "KEEPER: TreVeyon Henderson"
    assert tables["Top 300"].loc[32, "BLKWDW Pick"] == "KEEPER: Ashton Jeanty"
    assert tables["Top 300"].loc[176, "BLKWDW Pick"] == "KEEPER: TreVeyon Henderson"


def test_csv_xlsx_and_pdf_exports_are_complete_and_printable(tmp_path):
    tables = build_export_tables(
        _rankings(),
        pd.DataFrame(columns=["player_name", "owner_team", "keeper_round", "keeper_type"]),
        get_league("Drunk Sundays"),
    )

    csv_paths = write_csv_tables(tables, tmp_path / "csv")
    xlsx_path = write_excel_workbook(tables, tmp_path / "edgeiq.xlsx")
    pdf_path = write_pdf_draft_book(tables, tmp_path / "edgeiq.pdf")

    assert [path.name for path in csv_paths] == [
        "top_300.csv", "qb.csv", "rb.csv", "wr.csv", "te.csv", "k.csv", "def.csv", "blk_wdw_picks.csv"
    ]
    assert len(pd.read_csv(csv_paths[0])) == 300
    workbook = load_workbook(xlsx_path, read_only=False)
    assert workbook.sheetnames == list(tables)
    overall = workbook["Top 300"]
    assert overall.freeze_panes == "A2"
    assert overall.auto_filter.ref == overall.dimensions
    assert overall.sheet_properties.pageSetUpPr.fitToPage is True
    assert overall.page_setup.fitToWidth == 1
    assert overall.print_title_rows == "$1:$1"
    pdf_bytes = pdf_path.read_bytes()
    assert pdf_bytes.startswith(b"%PDF-1.4")
    assert pdf_bytes.endswith(b"%%EOF\n")
    assert pdf_bytes.count(b"/Type /Page ") >= 8
    assert b"EDGEIQ EMERGENCY DRUNK SUNDAYS DRAFT BOOK" in pdf_bytes
    assert b"Questionable - Knee" in pdf_bytes


def test_package_builder_uses_existing_rankings_and_writes_manifest(tmp_path):
    rankings = _rankings()
    keepers = pd.DataFrame(
        [{"owner_team": "Other", "player_name": "Player 1", "keeper_type": "rookie", "keeper_round": 3}]
    )
    calls = []

    outputs = build_draft_package(
        output_dir=tmp_path,
        ranking_builder=lambda league_key: calls.append(league_key) or rankings,
        keeper_loader=lambda league_name: calls.append(league_name) or keepers,
    )

    assert calls == ["drunk_sundays", "Drunk Sundays"]
    assert outputs["xlsx"].name == "edgeiq_drunk_sundays_emergency_draft_book.xlsx"
    assert outputs["pdf"].name == "edgeiq_drunk_sundays_emergency_draft_book.pdf"
    assert len(outputs["csv"]) == 8
    manifest = outputs["manifest"].read_text(encoding="utf-8")
    assert "Ranking source: build_draft_rankings('drunk_sundays')" in manifest
    assert "Keeper players remain ranked but are marked UNAVAILABLE" in manifest
