"""Print-ready XLSX output for EdgeIQ draft-book tables."""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _cell_value(value):
    return None if pd.isna(value) else value


def write_excel_workbook(tables, output_path):
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="17365D")
    keeper_fill = PatternFill("solid", fgColor="F4CCCC")
    pick_fill = PatternFill("solid", fgColor="FFF2CC")
    for name, table in tables.items():
        sheet = workbook.create_sheet(name)
        sheet.append(list(table.columns))
        for row in table.itertuples(index=False, name=None):
            sheet.append([_cell_value(value) for value in row])
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        headers = {cell.value: cell.column for cell in sheet[1]}
        for row_number in range(2, sheet.max_row + 1):
            if headers.get("Availability") and sheet.cell(row_number, headers["Availability"]).value == "UNAVAILABLE - KEEPER":
                for cell in sheet[row_number]:
                    cell.fill = keeper_fill
            elif headers.get("BLKWDW Pick") and sheet.cell(row_number, headers["BLKWDW Pick"]).value:
                for cell in sheet[row_number]:
                    cell.fill = pick_fill
        for column_number, column_name in enumerate(table.columns, 1):
            sample = [str(column_name)] + [str(value) for value in table.iloc[:100, column_number - 1].dropna()]
            sheet.column_dimensions[get_column_letter(column_number)].width = min(max(map(len, sample)) + 2, 34)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_LETTER
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = "1:1"
        sheet.oddFooter.center.text = "EdgeIQ Emergency Draft Book - &A - Page &P of &N"
    workbook.save(output_path)
    return output_path
